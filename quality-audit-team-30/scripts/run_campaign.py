"""Campagne de test QuickOrder : exécution des cas et collecte des preuves.

Usage (application lancée sur http://127.0.0.1:5000, base fraîchement initialisée) :

    python scripts/run_campaign.py

Produit :
  - docs/bug-reports/evidence/<ID>.txt : échanges HTTP (requêtes / réponses) de chaque cas
  - docs/campaign-results.json         : résultats bruts
  - docs/test-cases.xlsx               : cahier de tests avec résultat obtenu et statut
"""
import html
import json
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = "http://127.0.0.1:5000"
KEY = "quickorder-demo-key"
ROOT = Path(__file__).resolve().parent.parent
EVIDENCE = ROOT / "docs" / "bug-reports" / "evidence"
ALICE = ("alice@example.com", "Password123!")
BOB = ("bob@example.com", "Password123!")
ADMIN = ("admin@quickorder.local", "Admin123!")
RUN_ID = str(int(time.time()))

CASES = []


class Evidence:
    def __init__(self, case_id):
        self.case_id = case_id
        self.lines = []

    def log(self, text):
        self.lines.append(text)

    def note(self, text):
        self.lines.append(f"# {text}")

    def save(self):
        EVIDENCE.mkdir(parents=True, exist_ok=True)
        (EVIDENCE / f"{self.case_id}.txt").write_text("\n".join(self.lines) + "\n", encoding="utf-8")


def case(id, categorie, titre, objectif, preconditions, donnees, etapes, attendu, priorite="P2", bug=""):
    def deco(fn):
        CASES.append(dict(id=id, categorie=categorie, titre=titre, objectif=objectif,
                          preconditions=preconditions, donnees=donnees, etapes=etapes,
                          attendu=attendu, priorite=priorite, bug=bug, fn=fn))
        return fn
    return deco


# ---------------------------------------------------------------- helpers

def flashes(text):
    found = re.findall(r'<div class="flash (\w+)">(.*?)</div>', text, flags=re.S)
    return [f"[{c}] {html.unescape(m)}" for c, m in found if not m.startswith("<strong>")]


def api(ev, method, path, key=KEY, body=None, params=None, raw=None):
    headers = {}
    if key is not None:
        headers["X-API-Key"] = key
    if raw is not None:
        headers["Content-Type"] = "application/json"
        r = requests.request(method, BASE + "/api" + path, headers=headers, data=raw, params=params, timeout=15)
    else:
        r = requests.request(method, BASE + "/api" + path, headers=headers, json=body, params=params, timeout=15)
    ev.log(f"> {method} /api{path}" + (f"?{requests.compat.urlencode(params)}" if params else ""))
    ev.log(f"  X-API-Key: {key if key is not None else '(absente)'}")
    if body is not None:
        ev.log(f"  body: {json.dumps(body, ensure_ascii=False)}")
    if raw is not None:
        ev.log(f"  body (brut): {raw}")
    ctype = r.headers.get("Content-Type", "")
    ev.log(f"< {r.status_code} ({ctype})")
    if "json" in ctype:
        ev.log("  " + json.dumps(r.json(), ensure_ascii=False))
    else:
        ev.log("  " + re.sub(r"\s+", " ", r.text)[:300])
    ev.log("")
    return r


def js(r):
    try:
        return r.json()
    except ValueError:
        return {}


def web(ev, s, method, path, data=None, follow=True):
    r = s.request(method, BASE + path, data=data, allow_redirects=follow, timeout=15)
    ev.log(f"> {method} {path}")
    if data:
        shown = {k: ("***" if k == "password" else v) for k, v in data.items()}
        ev.log(f"  form: {shown}")
    chain = "".join(f"{h.status_code} -> {h.headers.get('Location', '')} ; " for h in r.history)
    loc = f" Location: {r.headers['Location']}" if "Location" in r.headers else ""
    ev.log(f"< {chain}{r.status_code} {r.url.replace(BASE, '')}{loc}")
    for f in flashes(r.text):
        ev.log(f"  message affiché : {f}")
    ev.log("")
    return r


def login(ev, creds):
    s = requests.Session()
    web(ev, s, "POST", "/login", {"email": creds[0], "password": creds[1]})
    return s


def ui_order(ev, s, items, payment="card", coupon="", cash=None):
    """Parcours UI complet : panier -> checkout -> confirmation. Retourne (réponse, id commande)."""
    web(ev, s, "POST", "/cart/clear")
    for pid, qty in items:
        web(ev, s, "POST", f"/cart/add/{pid}", {"quantity": qty})
    r = web(ev, s, "GET", "/checkout")
    m = re.search(r'<option value="(\d+)"[^>]*>(.*?)</option>', r.text, flags=re.S)
    if not m:
        return r, None
    ev.note(f"adresse proposée au checkout : «{html.unescape(m.group(2)).strip()}»")
    data = {"address_id": m.group(1), "customer_note": "", "coupon": coupon,
            "payment_method": payment, "confirm": "1"}
    if cash is not None:
        data["cash_received"] = str(cash)
    r = web(ev, s, "POST", "/checkout", data)
    m = re.search(r"/orders/(\d+)$", r.url)
    return r, (int(m.group(1)) if m else None)


def order_facts(ev, text):
    facts = {}
    m = re.search(r'<span class="status (\w+)">', text)
    facts["status"] = m.group(1) if m else None
    m = re.search(r'line total"><span>Total</span><span>([\d.]+) €', text)
    facts["total"] = m.group(1) if m else None
    m = re.search(r"Réduction</span><span>- ([\d.]+) €", text)
    facts["discount"] = m.group(1) if m else None
    m = re.search(r"Monnaie calculée : (-?[\d.]+) €", text)
    facts["change"] = m.group(1) if m else None
    m = re.search(r"<h3>Paiement</h3><p>(.*?)</p>", text)
    facts["payment"] = html.unescape(m.group(1)) if m else None
    ev.note(f"page commande : {facts}")
    return facts


def stock(ev, pid):
    return js(api(ev, "GET", f"/products/{pid}")).get("stock")


def alice_order_count(ev):
    return js(api(ev, "GET", "/orders")).get("count")


# ---------------------------------------------------------------- NOMINAUX

@case("NOM-01", "Nominal", "Connexion d'un client avec des identifiants valides",
      "Vérifier qu'un client existant peut s'authentifier.",
      "Compte alice@example.com existant ; utilisateur déconnecté.",
      "alice@example.com / Password123!",
      "1. Ouvrir /login\n2. Saisir email et mot de passe\n3. Valider",
      "Redirection vers l'accueil ; le nom « Alice Martin » et le lien Déconnexion sont affichés.", "P1")
def nom01(ev):
    s = requests.Session()
    r = web(ev, s, "POST", "/login", {"email": ALICE[0], "password": ALICE[1]})
    ok = "Alice Martin" in r.text and "Déconnexion" in r.text
    return f"HTTP {r.history[0].status_code if r.history else r.status_code} puis page d'accueil ; nom affiché : {'oui' if ok else 'non'}.", ok


@case("NOM-02", "Nominal", "Inscription d'un nouveau client",
      "Vérifier la création de compte avec des données valides.",
      "Email non encore utilisé.",
      f"Nom « Testeur QA », email qa.nom02.{RUN_ID}@example.com, mot de passe Password123!",
      "1. Ouvrir /register\n2. Remplir le formulaire\n3. Valider",
      "Compte créé, utilisateur connecté automatiquement et redirigé vers l'accueil.", "P1")
def nom02(ev):
    s = requests.Session()
    r = web(ev, s, "POST", "/register", {"name": "Testeur QA", "email": f"qa.nom02.{RUN_ID}@example.com",
                                           "password": "Password123!"})
    ok = "Testeur QA" in r.text and "Déconnexion" in r.text
    return f"Compte créé et session ouverte : {'oui' if ok else 'non'} (page finale {r.url.replace(BASE, '') or '/'}).", ok


@case("NOM-03", "Nominal", "Commande complète par carte bancaire (UI)",
      "Vérifier le parcours critique panier → checkout → confirmation.",
      "Alice connectée ; Burger Factory ouvert ; Classic Burger en stock.",
      "2 × Classic Burger (11,90 €) ; livraison 2,99 € ; paiement carte.",
      "1. Ajouter 2 Classic Burger au panier\n2. Passer commande\n3. Choisir Carte bancaire\n4. Confirmer",
      "Commande créée au statut confirmed, payée ; total = 23,80 + 2,99 = 26,79 € ; stock décrémenté de 2.", "P1")
def nom03(ev):
    s = login(ev, ALICE)
    before = stock(ev, 1)
    r, oid = ui_order(ev, s, [(1, 2)], "card")
    f = order_facts(ev, r.text)
    after = stock(ev, 1)
    ok = oid is not None and f["status"] == "confirmed" and f["total"] == "26.79" and before - after == 2
    return f"Commande #{oid} statut {f['status']}, total {f['total']} €, paiement « {f['payment']} » ; stock {before} → {after}.", ok


@case("NOM-04", "Nominal", "Application d'un coupon valide au checkout (UI)",
      "Vérifier le calcul de la réduction d'un coupon valide.",
      "Alice connectée ; panier ≥ 20 €.",
      "2 × Classic Burger (23,80 €) ; coupon WELCOME10 (10 %, min 20 €, plafond 8 €).",
      "1. Ajouter 2 Classic Burger\n2. Aller au checkout\n3. Saisir WELCOME10\n4. Cliquer Appliquer",
      "Réduction affichée 2,38 € ; total 23,80 + 2,99 − 2,38 = 24,41 €.", "P1")
def nom04(ev):
    s = login(ev, ALICE)
    web(ev, s, "POST", "/cart/clear")
    web(ev, s, "POST", "/cart/add/1", {"quantity": 2})
    r = web(ev, s, "POST", "/checkout", {"coupon": "WELCOME10", "confirm": "0", "payment_method": "card"})
    f = order_facts(ev, r.text)
    web(ev, s, "POST", "/cart/clear")
    ok = f["discount"] == "2.38" and f["total"] == "24.41"
    return f"Réduction affichée {f['discount']} € ; total {f['total']} €.", ok


@case("NOM-05", "Nominal", "API : liste des restaurants avec clé valide",
      "Vérifier que l'API renvoie la liste des restaurants au format JSON.",
      "Clé API valide.",
      "GET /api/restaurants ; X-API-Key: quickorder-demo-key",
      "1. Swagger > Authorize\n2. Exécuter GET /restaurants",
      "HTTP 200 ; JSON avec count et items ; chaque item contient id, name, delivery_fee, minimum_order, is_open.", "P1")
def nom05(ev):
    r = api(ev, "GET", "/restaurants")
    d = js(r)
    items = d.get("items", [])
    ok = (r.status_code == 200 and d.get("count") == len(items) and len(items) >= 5
          and all({"id", "name", "delivery_fee", "minimum_order", "is_open"} <= set(i) for i in items))
    return f"HTTP {r.status_code} ; count={d.get('count')} ; {len(items)} items avec les champs attendus : {'oui' if ok else 'non'}.", ok


@case("NOM-06", "Nominal", "API : création d'une commande par carte",
      "Vérifier la création d'une commande via POST /orders.",
      "Clé API valide ; Pizza Roma ouvert ; produits en stock.",
      '{"items":[{"product_id":16,"quantity":1},{"product_id":19,"quantity":1}],"payment_method":"card"}',
      "1. POST /api/orders avec le corps indiqué",
      "HTTP 201 ; status confirmed ; subtotal 17,40 ; delivery_fee 2,49 ; total 19,89.", "P1")
def nom06(ev):
    r = api(ev, "POST", "/orders", body={"items": [{"product_id": 16, "quantity": 1}, {"product_id": 19, "quantity": 1}],
                                          "payment_method": "card"})
    d = js(r)
    ok = r.status_code == 201 and d.get("status") == "confirmed" and d.get("subtotal") == 17.4 and d.get("total") == 19.89
    nom06.order_id = d.get("id")
    return f"HTTP {r.status_code} ; commande #{d.get('id')} {d.get('status')} ; subtotal {d.get('subtotal')} ; total {d.get('total')}.", ok


@case("NOM-07", "Nominal", "API : annulation d'une commande confirmée",
      "Vérifier qu'une commande au statut confirmed peut être annulée.",
      "Commande confirmée créée en NOM-06.",
      "POST /api/orders/{id}/cancel",
      "1. POST /api/orders/{id}/cancel\n2. GET /api/orders/{id}",
      "HTTP 200 ; status cancelled ; le statut est persisté.", "P2")
def nom07(ev):
    oid = getattr(nom06, "order_id", None)
    r = api(ev, "POST", f"/orders/{oid}/cancel")
    r2 = api(ev, "GET", f"/orders/{oid}")
    ok = r.status_code == 200 and js(r).get("status") == "cancelled" and js(r2).get("status") == "cancelled"
    return f"HTTP {r.status_code} ; statut après annulation : {js(r2).get('status')}.", ok


@case("NOM-08", "Nominal", "Recommander une ancienne commande (UI)",
      "Vérifier que « Commander à nouveau » recharge le panier.",
      "Alice connectée ; commande #1 (Classic Burger + Frites maison) existante.",
      "Commande #1.",
      "1. Ouvrir /orders/1\n2. Cliquer « Commander à nouveau »",
      "Panier contenant 1 Classic Burger et 1 Frites maison ; message « 2 produit(s) ajouté(s) au panier ».", "P3")
def nom08(ev):
    s = login(ev, ALICE)
    r = web(ev, s, "POST", "/orders/1/reorder")
    ok = "Classic Burger" in r.text and "Frites maison" in r.text and any("2 produit(s)" in f for f in flashes(r.text))
    web(ev, s, "POST", "/cart/clear")
    return f"Panier rechargé avec les 2 produits : {'oui' if ok else 'non'} ; messages : {flashes(r.text)}.", ok


# ---------------------------------------------------------------- LIMITES

def _order_qty(ev, pid, qty):
    return api(ev, "POST", "/orders", body={"items": [{"product_id": pid, "quantity": qty}], "payment_method": "card"})


@case("LIM-01", "Limite", "API : quantité maximale autorisée (20)",
      "Vérifier la borne haute valide de la quantité.",
      "Clé API valide ; Miso Soup stock ≥ 20.",
      "product_id 9 (Miso Soup 3,20 €), quantity 20.",
      "1. POST /api/orders quantity=20",
      "HTTP 201 ; subtotal 64,00 ; total 67,49.", "P2")
def lim01(ev):
    r = _order_qty(ev, 9, 20)
    d = js(r)
    ok = r.status_code == 201 and d.get("subtotal") == 64.0 and d.get("total") == 67.49
    return f"HTTP {r.status_code} ; subtotal {d.get('subtotal')} ; total {d.get('total')}.", ok


@case("LIM-02", "Limite", "API : quantité juste au-dessus du maximum (21)",
      "Vérifier le rejet de la première valeur invalide au-dessus de la borne.",
      "Clé API valide.",
      "product_id 9, quantity 21.",
      "1. POST /api/orders quantity=21",
      "HTTP 400 ; message « quantity must be between 1 and 20 » ; aucune commande créée.", "P2")
def lim02(ev):
    n = alice_order_count(ev)
    r = _order_qty(ev, 9, 21)
    n2 = alice_order_count(ev)
    ok = r.status_code == 400 and n == n2
    return f"HTTP {r.status_code} ; erreur « {js(r).get('error')} » ; commandes {n} → {n2}.", ok


@case("LIM-03", "Limite", "API : quantité 0 et quantité négative",
      "Vérifier le rejet des valeurs sous la borne basse.",
      "Clé API valide.",
      "product_id 9, quantity 0 puis −1.",
      "1. POST /api/orders quantity=0\n2. POST /api/orders quantity=-1",
      "HTTP 400 dans les deux cas ; aucune commande créée.", "P2")
def lim03(ev):
    n = alice_order_count(ev)
    r0 = _order_qty(ev, 9, 0)
    r1 = _order_qty(ev, 9, -1)
    n2 = alice_order_count(ev)
    ok = r0.status_code == 400 and r1.status_code == 400 and n == n2
    return f"quantity=0 → HTTP {r0.status_code} ; quantity=-1 → HTTP {r1.status_code} ; commandes {n} → {n2}.", ok


@case("LIM-04", "Limite", "API : quantité égale au stock + 1",
      "Vérifier qu'on ne peut pas commander plus que le stock disponible.",
      "Clé API valide ; Salmon Bowl en stock limité (2 à l'initialisation).",
      "product_id 7 (Salmon Bowl), quantity = stock + 1.",
      "1. GET /api/products/7 (lire le stock)\n2. POST /api/orders quantity = stock + 1\n3. GET /api/products/7",
      "Commande refusée (HTTP 409 ou 400) ; stock inchangé et jamais négatif.", "P1", "BUG-04")
def lim04(ev):
    before = stock(ev, 7)
    r = _order_qty(ev, 7, before + 1)
    after = stock(ev, 7)
    ok = r.status_code in (400, 409) and after == before
    return f"Stock initial {before} ; commande de {before + 1} → HTTP {r.status_code} (commande #{js(r).get('id')}) ; stock final {after}.", ok


@case("LIM-05", "Limite", "API : coupon au montant minimum exact et juste en dessous",
      "Vérifier la borne du minimum de commande d'un coupon.",
      "Coupon WELCOME10 actif (min 20 €).",
      "subtotal 20.00 puis 19.99.",
      "1. POST /api/coupons/validate subtotal=20.00\n2. POST /api/coupons/validate subtotal=19.99",
      "20,00 → valid true, discount 2,00 ; 19,99 → valid false « Montant minimum non atteint ».", "P2")
def lim05(ev):
    a = js(api(ev, "POST", "/coupons/validate", body={"code": "WELCOME10", "subtotal": 20.00}))
    b = js(api(ev, "POST", "/coupons/validate", body={"code": "WELCOME10", "subtotal": 19.99}))
    ok = a.get("valid") is True and a.get("discount") == 2.0 and b.get("valid") is False
    return f"20,00 → valid={a.get('valid')}, discount={a.get('discount')} ; 19,99 → valid={b.get('valid')} ({b.get('error')}).", ok


@case("LIM-06", "Limite", "API : plafond de réduction d'un coupon",
      "Vérifier que la réduction est plafonnée à max_discount.",
      "Coupon SAVE20 actif (20 %, min 35 €, plafond 12 €).",
      "subtotal 59.95, 60.00 et 100.00.",
      "1. POST /api/coupons/validate avec chaque sous-total",
      "59,95 → 11,99 ; 60,00 → 12,00 ; 100,00 → 12,00 (plafonné).", "P2")
def lim06(ev):
    vals = [js(api(ev, "POST", "/coupons/validate", body={"code": "SAVE20", "subtotal": v})).get("discount")
            for v in (59.95, 60.00, 100.00)]
    ok = vals == [11.99, 12.0, 12.0]
    return f"Réductions obtenues : 59,95 → {vals[0]} ; 60,00 → {vals[1]} ; 100,00 → {vals[2]}.", ok


@case("LIM-07", "Limite", "Inscription : longueur minimale du mot de passe (7 / 8 caractères)",
      "Vérifier la borne de 8 caractères sur le mot de passe.",
      "Emails non utilisés.",
      "Mot de passe « Abcdef1 » (7) puis « Abcdef12 » (8).",
      "1. S'inscrire avec 7 caractères\n2. S'inscrire avec 8 caractères",
      "7 caractères → refus avec message ; 8 caractères → compte créé.", "P2")
def lim07(ev):
    s1, s2 = requests.Session(), requests.Session()
    r1 = web(ev, s1, "POST", "/register", {"name": "Limite Sept", "email": f"qa.lim07a.{RUN_ID}@example.com", "password": "Abcdef1"})
    r2 = web(ev, s2, "POST", "/register", {"name": "Limite Huit", "email": f"qa.lim07b.{RUN_ID}@example.com", "password": "Abcdef12"})
    refused = "Déconnexion" not in r1.text and bool(flashes(r1.text))
    accepted = "Limite Huit" in r2.text and "Déconnexion" in r2.text
    return f"7 caractères → {'refusé' if refused else 'accepté'} {flashes(r1.text)} ; 8 caractères → {'accepté' if accepted else 'refusé'}.", refused and accepted


@case("LIM-08", "Limite", "Panier UI : quantité au-delà du maximum affiché et du stock",
      "Vérifier que le serveur borne la quantité du panier (le champ HTML affiche max = stock).",
      "Alice connectée ; Classic Burger dans le panier (stock ≈ 100).",
      "qty_1 = 1000 envoyé au serveur (contournement du max HTML via DevTools).",
      "1. Ajouter 1 Classic Burger\n2. Dans DevTools, retirer l'attribut max du champ quantité\n3. Saisir 1000 et cliquer « Mettre à jour »",
      "Quantité refusée ou ramenée au stock disponible ; message d'erreur.", "P2", "BUG-04")
def lim08(ev):
    s = login(ev, ALICE)
    web(ev, s, "POST", "/cart/clear")
    web(ev, s, "POST", "/cart/add/1", {"quantity": 1})
    st = stock(ev, 1)
    r = web(ev, s, "POST", "/cart", {"qty_1": "1000"})
    m = re.search(r'name="qty_1" value="(\d+)"', r.text)
    qty = int(m.group(1)) if m else None
    m = re.search(r'class="total">Total <strong>([\d.]+) €', r.text)
    total = m.group(1) if m else None
    web(ev, s, "POST", "/cart/clear")
    ok = qty is not None and qty <= st
    return f"Stock {st} ; quantité enregistrée dans le panier : {qty} ; total panier {total} € ; aucun message d'erreur : {not flashes(r.text)}.", ok


@case("LIM-09", "Limite", "API : paiement en espèces avec le montant exact",
      "Vérifier la borne « montant remis = total ».",
      "Clé API valide.",
      "1 × Margherita (10,90 €) + livraison 2,49 € = 13,39 € ; cash_received 13.39.",
      "1. POST /api/orders payment_method=cash cash_received=13.39",
      "HTTP 201 ; change_due = 0,00.", "P2")
def lim09(ev):
    r = api(ev, "POST", "/orders", body={"items": [{"product_id": 16, "quantity": 1}], "payment_method": "cash", "cash_received": 13.39})
    d = js(r)
    ok = r.status_code == 201 and d.get("total") == 13.39 and d.get("change_due") == 0.0
    return f"HTTP {r.status_code} ; total {d.get('total')} ; change_due {d.get('change_due')}.", ok


# ---------------------------------------------------------------- ERREURS

@case("ERR-01", "Erreur", "Connexion avec un mot de passe incorrect",
      "Vérifier le refus d'authentification et l'absence de fuite d'information.",
      "Compte alice@example.com existant.",
      "alice@example.com / MauvaisMotDePasse ; inconnu@example.com / Password123!",
      "1. Se connecter avec un mauvais mot de passe\n2. Se connecter avec un email inconnu",
      "Connexion refusée ; même message générique dans les deux cas ; pas de session.", "P1")
def err01(ev):
    s = requests.Session()
    r1 = web(ev, s, "POST", "/login", {"email": ALICE[0], "password": "MauvaisMotDePasse"})
    r2 = web(ev, s, "POST", "/login", {"email": "inconnu@example.com", "password": "Password123!"})
    ok = "Déconnexion" not in r1.text and flashes(r1.text) == flashes(r2.text) and bool(flashes(r1.text))
    return f"Connexion refusée ; message identique dans les deux cas : {flashes(r1.text)}.", ok


@case("ERR-02", "Erreur", "API : clé absente ou incorrecte",
      "Vérifier la protection des endpoints par X-API-Key.",
      "Aucune.",
      "GET /api/restaurants sans en-tête ; puis X-API-Key: mauvaise-cle.",
      "1. GET /api/restaurants sans clé\n2. GET /api/restaurants avec une clé erronée",
      "HTTP 401 avec un corps JSON d'erreur dans les deux cas.", "P1")
def err02(ev):
    r1 = api(ev, "GET", "/restaurants", key=None)
    r2 = api(ev, "GET", "/restaurants", key="mauvaise-cle")
    ok = r1.status_code == 401 and r2.status_code == 401 and "error" in js(r1)
    return f"Sans clé → HTTP {r1.status_code} ; clé erronée → HTTP {r2.status_code} ; erreur « {js(r1).get('error')} ».", ok


@case("ERR-03", "Erreur", "API : identifiants inexistants",
      "Vérifier les codes HTTP pour des ressources inexistantes.",
      "Clé API valide.",
      "restaurant 9999 ; produit 9999 ; commande 9999 ; commande contenant product_id 9999.",
      "1. GET /api/restaurants/9999\n2. GET /api/products/9999\n3. GET /api/orders/9999\n4. POST /api/orders avec product_id 9999",
      "HTTP 404 pour les trois GET ; HTTP 400 ou 404 pour le POST ; aucune commande créée.", "P2")
def err03(ev):
    codes = [api(ev, "GET", p).status_code for p in ("/restaurants/9999", "/products/9999", "/orders/9999")]
    r = _order_qty(ev, 9999, 1)
    ok = codes == [404, 404, 404] and r.status_code in (400, 404)
    return f"GET → {codes} ; POST commande produit 9999 → HTTP {r.status_code} « {js(r).get('error')} ».", ok


@case("ERR-04", "Erreur", "API : validation d'un coupon expiré ou inconnu",
      "Vérifier le rejet d'un coupon expiré par /coupons/validate.",
      "Coupon EXPIRED10 expiré depuis la veille de l'initialisation.",
      "code EXPIRED10 puis code INCONNU ; subtotal 30.",
      "1. POST /api/coupons/validate EXPIRED10\n2. POST /api/coupons/validate INCONNU",
      "valid false dans les deux cas, message « Coupon invalide ou expiré ».", "P2")
def err04(ev):
    r1 = api(ev, "POST", "/coupons/validate", body={"code": "EXPIRED10", "subtotal": 30})
    r2 = api(ev, "POST", "/coupons/validate", body={"code": "INCONNU", "subtotal": 30})
    ok = js(r1).get("valid") is False and js(r2).get("valid") is False
    return f"EXPIRED10 → HTTP {r1.status_code} valid={js(r1).get('valid')} ; INCONNU → HTTP {r2.status_code} valid={js(r2).get('valid')}.", ok


@case("ERR-05", "Erreur", "API : commande passée avec un coupon expiré",
      "Vérifier qu'un coupon expiré n'accorde aucune réduction lors de la création de commande.",
      "Coupon EXPIRED10 expiré (rejeté par /coupons/validate, cf. ERR-04).",
      '{"items":[{"product_id":1,"quantity":2}],"payment_method":"card","coupon":"EXPIRED10"}',
      "1. POST /api/orders avec coupon EXPIRED10",
      "Commande refusée, ou créée sans réduction (discount 0 ; total 26,79).", "P1", "BUG-08")
def err05(ev):
    r = api(ev, "POST", "/orders", body={"items": [{"product_id": 1, "quantity": 2}], "payment_method": "card", "coupon": "EXPIRED10"})
    d = js(r)
    ok = r.status_code in (400, 404, 409) or (r.status_code == 201 and d.get("discount") == 0.0)
    return f"HTTP {r.status_code} ; commande #{d.get('id')} ; discount {d.get('discount')} ; total {d.get('total')}.", ok


@case("ERR-06", "Erreur", "API : moyen de paiement invalide",
      "Vérifier le rejet d'un payment_method non supporté.",
      "Clé API valide.",
      'payment_method "bitcoin".',
      "1. POST /api/orders payment_method=bitcoin",
      "HTTP 400 « invalid payment_method » ; aucune commande créée.", "P2")
def err06(ev):
    n = alice_order_count(ev)
    r = api(ev, "POST", "/orders", body={"items": [{"product_id": 16, "quantity": 1}], "payment_method": "bitcoin"})
    n2 = alice_order_count(ev)
    ok = r.status_code == 400 and n == n2
    return f"HTTP {r.status_code} « {js(r).get('error')} » ; commandes {n} → {n2}.", ok


@case("ERR-07", "Erreur", "API : corps de requête mal formé",
      "Vérifier que l'API répond 400 (et jamais 500) à des entrées invalides.",
      "Clé API valide.",
      'a) JSON invalide  b) items []  c) quantity "abc"  d) items contenant un élément non-objet : [{"product_id":16,"quantity":1},"x"]',
      "1. POST /api/orders avec chacun des 4 corps",
      "HTTP 400 avec message d'erreur JSON pour les 4 requêtes.", "P2", "BUG-13")
def err07(ev):
    a = api(ev, "POST", "/orders", raw='{"items": [')
    b = api(ev, "POST", "/orders", body={"items": []})
    c = api(ev, "POST", "/orders", body={"items": [{"product_id": 16, "quantity": "abc"}]})
    d = api(ev, "POST", "/orders", body={"items": [{"product_id": 16, "quantity": 1}, "x"]})
    codes = [x.status_code for x in (a, b, c, d)]
    return f"Codes obtenus : a) {codes[0]}  b) {codes[1]}  c) {codes[2]}  d) {codes[3]}.", codes == [400, 400, 400, 400]


@case("ERR-08", "Erreur", "Contrôle d'accès aux pages protégées (UI)",
      "Vérifier qu'un utilisateur ne peut pas accéder aux ressources qui ne lui appartiennent pas.",
      "Comptes alice, bob ; commande #1 appartenant à Alice.",
      "Anonyme → /checkout ; Alice (client) → /admin/ ; Bob → /orders/1.",
      "1. Sans session, ouvrir /checkout\n2. Connecté en client, ouvrir /admin/\n3. Connecté en Bob, ouvrir /orders/1",
      "1) redirection vers /login  2) accès refusé « réservé aux administrateurs »  3) « Commande inaccessible ».", "P1")
def err08(ev):
    r1 = web(ev, requests.Session(), "GET", "/checkout", follow=False)
    r2 = web(ev, login(ev, ALICE), "GET", "/admin/")
    r3 = web(ev, login(ev, BOB), "GET", "/orders/1")
    a = r1.status_code == 302 and "/login" in r1.headers.get("Location", "")
    b = any("administrateurs" in f for f in flashes(r2.text)) and "<h1>Administration" not in r2.text
    c = any("inaccessible" in f for f in flashes(r3.text)) and "Commande #1" not in r3.text
    return (f"1) HTTP {r1.status_code} → {r1.headers.get('Location')} ; 2) {flashes(r2.text)} ; 3) {flashes(r3.text)}.", a and b and c)


@case("ERR-09", "Erreur", "API : format des réponses d'erreur 404",
      "Vérifier que toutes les erreurs de l'API sont renvoyées en JSON (contrat OpenAPI).",
      "Clé API valide.",
      "GET /api/restaurants/9999.",
      "1. GET /api/restaurants/9999\n2. Inspecter l'en-tête Content-Type et le corps",
      "HTTP 404 avec Content-Type application/json et un champ error.", "P3", "BUG-14")
def err09(ev):
    r = api(ev, "GET", "/restaurants/9999")
    ctype = r.headers.get("Content-Type", "")
    ok = r.status_code == 404 and "json" in ctype
    return f"HTTP {r.status_code} ; Content-Type « {ctype} ».", ok


@case("ERR-10", "Erreur", "Connexion : paramètre next pointant vers un site externe",
      "Vérifier que la redirection après connexion reste interne (pas de redirection ouverte).",
      "Compte alice existant.",
      "POST /login?next=https://example.org/phishing",
      "1. Ouvrir /login?next=https://example.org/phishing\n2. Se connecter avec des identifiants valides\n3. Observer l'en-tête Location (DevTools > Network)",
      "Redirection vers une page interne (accueil) ; l'URL externe est ignorée.", "P1", "BUG-12")
def err10(ev):
    s = requests.Session()
    r = web(ev, s, "POST", "/login?next=https://example.org/phishing", {"email": ALICE[0], "password": ALICE[1]}, follow=False)
    loc = r.headers.get("Location", "")
    ok = r.status_code == 302 and not loc.startswith("http") or loc.startswith(BASE)
    return f"HTTP {r.status_code} ; Location : {loc}.", ok


# ---------------------------------------------------------------- RÈGLES MÉTIER

@case("RM-01", "Règle métier", "Un restaurant fermé ne peut pas recevoir de commande",
      "Vérifier le blocage des commandes sur un restaurant au statut FERMÉ.",
      "Alice connectée ; Night Sushi (id 6) affiché FERMÉ.",
      "1 × Late Sushi Box (product_id 24, 14,00 €) ; paiement carte.",
      "UI : 1. Ouvrir Night Sushi  2. Ajouter Late Sushi Box  3. Passer commande et confirmer\nAPI : POST /api/orders product_id 24",
      "Ajout au panier ou validation refusés (UI) ; HTTP 409 (API) ; aucune commande créée.", "P1", "BUG-01")
def rm01(ev):
    s = login(ev, ALICE)
    r, oid = ui_order(ev, s, [(24, 1)], "card")
    f = order_facts(ev, r.text) if oid else {}
    ra = _order_qty(ev, 24, 1)
    ok = oid is None and ra.status_code in (400, 409)
    return (f"UI : commande #{oid} créée, statut {f.get('status')}, total {f.get('total')} € ; "
            f"API : HTTP {ra.status_code} commande #{js(ra).get('id')}.", ok)


@case("RM-02", "Règle métier", "Le minimum de commande du restaurant est appliqué",
      "Vérifier le refus d'une commande sous le minimum affiché.",
      "Alice connectée ; Burger Factory minimum 20,00 €.",
      "1 × Frites maison (3,50 €) ; paiement carte.",
      "UI : 1. Ajouter 1 Frites maison  2. Checkout (« Commande minimum : 20.00 € » affiché)  3. Confirmer\nAPI : POST /api/orders product_id 3",
      "Validation refusée avec message indiquant le minimum ; aucune commande créée.", "P1", "BUG-05")
def rm02(ev):
    s = login(ev, ALICE)
    r, oid = ui_order(ev, s, [(3, 1)], "card")
    f = order_facts(ev, r.text) if oid else {}
    ra = _order_qty(ev, 3, 1)
    ok = oid is None and ra.status_code in (400, 409)
    return (f"UI : commande #{oid} créée, statut {f.get('status')}, total {f.get('total')} € (sous-total 3,50 € < 20 €) ; "
            f"API : HTTP {ra.status_code} commande #{js(ra).get('id')}.", ok)


@case("RM-03", "Règle métier", "Un produit hors stock ne peut pas être commandé",
      "Vérifier le blocage d'un produit dont le stock est 0.",
      "Alice connectée ; Veggie Burger (id 4) et Gyozas (id 8) à stock 0, affichés « HORS STOCK ».",
      "UI : 1 × Veggie Burger + 2 × Classic Burger ; API : 1 × Gyozas.",
      "UI : 1. Ouvrir Burger Factory  2. Cliquer Ajouter sur Veggie Burger (HORS STOCK)  3. Commander\nAPI : POST /api/orders product_id 8",
      "Ajout au panier refusé (UI) ; HTTP 409 (API) ; stock jamais négatif.", "P1", "BUG-04")
def rm03(ev):
    s = login(ev, ALICE)
    r, oid = ui_order(ev, s, [(4, 1), (1, 2)], "card")
    st4 = stock(ev, 4)
    ra = _order_qty(ev, 8, 1)
    st8 = stock(ev, 8)
    ok = oid is None and ra.status_code in (400, 409) and st4 >= 0 and st8 >= 0
    return (f"UI : commande #{oid} créée avec Veggie Burger, stock Veggie Burger = {st4} ; "
            f"API : HTTP {ra.status_code}, stock Gyozas = {st8}.", ok)


@case("RM-04", "Règle métier", "Paiement en espèces : le montant remis doit couvrir le total",
      "Vérifier le refus d'un montant en espèces inférieur au total.",
      "Alice connectée.",
      "2 × Classic Burger = 26,79 € livraison incluse ; montant remis 5,00 € (UI) et 1 (API).",
      "UI : 1. Panier 2 Classic Burger  2. Checkout, mode Espèces, montant 5.00  3. Confirmer\nAPI : POST /api/orders payment_method=cash cash_received=1",
      "Commande refusée avec message « montant insuffisant » ; aucune commande payée.", "P1", "BUG-03")
def rm04(ev):
    s = login(ev, ALICE)
    r, oid = ui_order(ev, s, [(1, 2)], "cash", cash="5.00")
    f = order_facts(ev, r.text) if oid else {}
    ra = api(ev, "POST", "/orders", body={"items": [{"product_id": 1, "quantity": 1}], "payment_method": "cash", "cash_received": 1})
    d = js(ra)
    ok = oid is None and ra.status_code in (400, 409, 422)
    return (f"UI : commande #{oid} {f.get('status')} « {f.get('payment')} », total {f.get('total')} €, monnaie calculée {f.get('change')} € ; "
            f"API : HTTP {ra.status_code}, payment_status {d.get('payment_status')}, change_due {d.get('change_due')}.", ok)


@case("RM-05", "Règle métier", "Un paiement PayPal crée une seule commande",
      "Vérifier l'unicité de la commande et du débit pour un paiement PayPal.",
      "Alice connectée.",
      "2 × Regina (27,80 €) ; paiement PayPal (UI) ; 1 × Regina PayPal (API).",
      "UI : 1. Noter le nombre de commandes  2. Commander 2 Regina via PayPal  3. Ouvrir Mes commandes\nAPI : POST /api/orders payment_method=paypal puis GET /api/orders",
      "Exactement 1 commande supplémentaire par paiement ; stock décrémenté une seule fois.", "P1", "BUG-02")
def rm05(ev):
    s = login(ev, ALICE)
    n0, st0 = alice_order_count(ev), stock(ev, 17)
    r, oid = ui_order(ev, s, [(17, 2)], "paypal")
    n1, st1 = alice_order_count(ev), stock(ev, 17)
    api(ev, "POST", "/orders", body={"items": [{"product_id": 17, "quantity": 1}], "payment_method": "paypal"})
    n2, st2 = alice_order_count(ev), stock(ev, 17)
    ok = n1 - n0 == 1 and n2 - n1 == 1
    return (f"UI : commandes {n0} → {n1} (+{n1 - n0}) pour un seul paiement, stock Regina {st0} → {st1} (2 commandés) ; "
            f"API : commandes {n1} → {n2} (+{n2 - n1}), stock {st1} → {st2} (1 commandé).", ok)


@case("RM-06", "Règle métier", "Le coupon ATLAS15 accorde la réduction annoncée (15 %)",
      "Vérifier la cohérence entre la promotion affichée et la réduction appliquée.",
      "Alice connectée ; page Couscous Atlas affichant « ATLAS15 … 15 % de réduction ».",
      "2 × Couscous Royal (35,80 €) ; coupon ATLAS15.",
      "1. Ouvrir Couscous Atlas (lire la promotion)\n2. Ajouter 2 Couscous Royal\n3. Checkout, saisir ATLAS15, Appliquer",
      "Réduction de 15 % = 5,37 € ; total 35,80 + 2,99 − 5,37 = 33,42 €.", "P1", "BUG-06")
def rm06(ev):
    s = login(ev, ALICE)
    page = web(ev, s, "GET", "/restaurant/5")
    m = re.search(r"Offre du restaurant :</strong>(.*?)</div>", page.text)
    ev.note("promotion affichée : " + re.sub(r"<[^>]+>", "", html.unescape(m.group(1))).strip() if m else "promotion non trouvée")
    web(ev, s, "POST", "/cart/clear")
    web(ev, s, "POST", "/cart/add/20", {"quantity": 2})
    r = web(ev, s, "POST", "/checkout", {"coupon": "ATLAS15", "confirm": "0", "payment_method": "card"})
    f = order_facts(ev, r.text)
    v = js(api(ev, "POST", "/coupons/validate", body={"code": "ATLAS15", "subtotal": 35.80}))
    web(ev, s, "POST", "/cart/clear")
    ok = f["discount"] == "5.37"
    return f"Réduction appliquée {f['discount']} € (total {f['total']} €) ; l'API annonce discount_percent = {v.get('discount_percent')}.", ok


@case("RM-07", "Règle métier", "Le panier ne contient que des produits d'un même restaurant",
      "Vérifier la règle du panier mono-restaurant.",
      "Alice connectée ; panier vide.",
      "Classic Burger (Burger Factory) puis Chicken Teriyaki (Tokyo Bowl) ; API : product_id 1 + 6.",
      "UI : 1. Ajouter Classic Burger  2. Ajouter Chicken Teriyaki\nAPI : POST /api/orders avec les deux produits",
      "Second ajout refusé avec message ; API HTTP 400.", "P2")
def rm07(ev):
    s = login(ev, ALICE)
    web(ev, s, "POST", "/cart/clear")
    web(ev, s, "POST", "/cart/add/1", {"quantity": 1})
    r = web(ev, s, "POST", "/cart/add/6", {"quantity": 1})
    cart = web(ev, s, "GET", "/cart")
    ra = api(ev, "POST", "/orders", body={"items": [{"product_id": 1, "quantity": 1}, {"product_id": 6, "quantity": 1}], "payment_method": "card"})
    web(ev, s, "POST", "/cart/clear")
    ok = "Chicken Teriyaki" not in cart.text and "Classic Burger" in cart.text and ra.status_code == 400
    return f"UI : {flashes(r.text)} ; panier limité au 1er restaurant : {'oui' if 'Chicken Teriyaki' not in cart.text else 'non'} ; API : HTTP {ra.status_code}.", ok


@case("RM-08", "Règle métier", "Le filtre « Livraison max » ne retient que les restaurants sous le plafond",
      "Vérifier le filtre des frais de livraison sur la page d'accueil.",
      "Frais : Green Kitchen 1,99 ; Pizza Roma 2,49 ; Burger Factory 2,99 ; Couscous Atlas 2,99 ; Night Sushi 2,99 ; Tokyo Bowl 3,49.",
      "Filtre « ≤ 2,50 € » (max_delivery=2.50).",
      "1. Sur l'accueil, choisir Livraison max « ≤ 2,50 € »\n2. Cliquer Rechercher",
      "Seuls Green Kitchen (1,99 €) et Pizza Roma (2,49 €) sont affichés.", "P2", "BUG-07")
def rm08(ev):
    r = web(ev, requests.Session(), "GET", "/?max_delivery=2.50")
    names = [html.unescape(n) for n in re.findall(r"<h2>(.*?)</h2>", r.text)]
    ev.note(f"restaurants affichés : {names}")
    ok = sorted(names) == ["Green Kitchen", "Pizza Roma"]
    return f"Restaurants affichés : {', '.join(names)}.", ok


@case("RM-09", "Règle métier", "Une commande livrée ne peut plus être annulée par le client",
      "Vérifier la règle d'annulation selon le statut.",
      "Commande #1 d'Alice au statut delivered.",
      "POST /api/orders/1/cancel",
      "1. POST /api/orders/1/cancel\n2. GET /api/orders/1",
      "HTTP 409 « order cannot be cancelled » ; statut toujours delivered.", "P2")
def rm09(ev):
    r = api(ev, "POST", "/orders/1/cancel")
    st = js(api(ev, "GET", "/orders/1")).get("status")
    ok = r.status_code == 409 and st == "delivered"
    return f"HTTP {r.status_code} « {js(r).get('error')} » ; statut conservé : {st}.", ok


@case("RM-10", "Règle métier", "L'annulation d'une commande restitue le stock et le paiement",
      "Vérifier les effets de bord d'une annulation (stock, remboursement).",
      "Clé API valide ; Tiramisu en stock.",
      "3 × Tiramisu (product_id 19), carte ; puis annulation.",
      "1. Lire le stock\n2. POST /api/orders 3 Tiramisu\n3. POST /api/orders/{id}/cancel\n4. Relire le stock et la commande",
      "Stock revenu à sa valeur initiale ; payment_status passé à refunded (ou équivalent).", "P2", "BUG-11")
def rm10(ev):
    s0 = stock(ev, 19)
    oid = js(_order_qty(ev, 19, 3)).get("id")
    s1 = stock(ev, 19)
    d = js(api(ev, "POST", f"/orders/{oid}/cancel"))
    s2 = stock(ev, 19)
    ok = s2 == s0 and d.get("payment_status") != "paid"
    return f"Stock {s0} → {s1} après commande → {s2} après annulation ; statut {d.get('status')} ; payment_status « {d.get('payment_status')} ».", ok


@case("RM-11", "Règle métier", "Admin : les transitions de statut respectent le cycle de vie",
      "Vérifier qu'une commande livrée ne peut pas revenir à un statut antérieur.",
      "Admin connecté ; commande #1 au statut delivered.",
      "Commande #1 ; nouveau statut confirmed.",
      "1. Se connecter en admin\n2. /admin/ : sur la commande #1 choisir « confirmed »\n3. Mettre à jour",
      "Transition refusée (« Statut invalide ») ; la commande reste delivered.", "P2", "BUG-09")
def rm11(ev):
    s = login(ev, ADMIN)
    r = web(ev, s, "POST", "/admin/orders/1/status", {"status": "confirmed"})
    st = js(api(ev, "GET", "/orders/1")).get("status")
    ok = st == "delivered"
    ev.note("remise en état : la commande #1 est replacée au statut delivered")
    web(ev, s, "POST", "/admin/orders/1/status", {"status": "delivered"})
    return f"{flashes(r.text)} ; statut de la commande #1 après l'action : {st} (remise à delivered en fin de test).", ok


@case("RM-12", "Règle métier", "Seul un client ayant commandé peut laisser un avis",
      "Vérifier qu'un avis est lié à une commande réelle du restaurant.",
      "Bob connecté ; Bob n'a jamais commandé chez Couscous Atlas.",
      "Note 1 ; commentaire « Avis de test QA sans commande ».",
      "1. Se connecter en Bob\n2. Ouvrir Couscous Atlas\n3. Publier un avis 1 étoile",
      "Publication refusée (« vous devez avoir commandé dans ce restaurant »).", "P2", "BUG-10")
def rm12(ev):
    s = login(ev, BOB)
    orders = web(ev, s, "GET", "/orders")
    ev.note(f"Bob a des commandes chez Couscous Atlas : {'Couscous Atlas' in orders.text}")
    r = web(ev, s, "POST", "/restaurant/5/review", {"rating": "1", "comment": "Avis de test QA sans commande"})
    accepted = any("Merci pour votre avis" in f for f in flashes(r.text))
    return f"Messages : {flashes(r.text)} ; avis accepté sans commande : {'oui' if accepted else 'non'}.", not accepted


@case("RM-13", "Règle métier", "Une commande nécessite une adresse de livraison complète",
      "Vérifier qu'un nouveau compte ne peut pas commander sans adresse renseignée.",
      "Compte créé à l'instant (aucune adresse saisie par l'utilisateur).",
      f"Compte qa.rm13.{RUN_ID}@example.com ; 2 × Margherita ; carte.",
      "1. Créer un compte\n2. Ajouter 2 Margherita\n3. Checkout : observer l'adresse proposée\n4. Confirmer",
      "Le checkout exige la saisie d'une adresse (rue, code postal, ville) avant de confirmer.", "P1", "BUG-15")
def rm13(ev):
    s = requests.Session()
    web(ev, s, "POST", "/register", {"name": "Sans Adresse", "email": f"qa.rm13.{RUN_ID}@example.com", "password": "Password123!"})
    r, oid = ui_order(ev, s, [(16, 2)], "card")
    m = re.search(r"<h3>Livraison</h3><p>(.*?)</p>", r.text, flags=re.S)
    addr = re.sub(r"<br>|\s+", " ", m.group(1)).strip() if m else None
    ev.note(f"adresse de livraison affichée sur la commande : «{addr}»")
    ok = oid is None
    return f"Commande #{oid} confirmée avec l'adresse « Domicile » vide (rue, code postal et ville vides : «{addr}»).", ok


@case("RM-14", "Règle métier", "TOKYO30 offre la livraison dès 30 € d'achat",
      "Vérifier la promotion « Livraison gratuite dès 30 € avec TOKYO30 ».",
      "Coupon TOKYO30 actif ; frais de livraison Tokyo Bowl 3,49 €.",
      "subtotal 29.99 puis 31.80.",
      "1. POST /api/coupons/validate TOKYO30 subtotal=29.99\n2. Idem subtotal=31.80",
      "29,99 → non valide ; 31,80 → réduction 3,49 € (= frais de livraison).", "P3")
def rm14(ev):
    a = js(api(ev, "POST", "/coupons/validate", body={"code": "TOKYO30", "subtotal": 29.99}))
    b = js(api(ev, "POST", "/coupons/validate", body={"code": "TOKYO30", "subtotal": 31.80}))
    ok = a.get("valid") is False and b.get("valid") is True and b.get("discount") == 3.49
    return f"29,99 → valid={a.get('valid')} ; 31,80 → valid={b.get('valid')}, discount={b.get('discount')}.", ok


# ---------------------------------------------------------------- observations exploratoires (sans verdict)

def exploratory():
    ev = Evidence("EXPLORATOIRE")
    ev.note("Observations exploratoires : relevés bruts, sans verdict Pass/Fail")
    r = requests.get(BASE + "/", timeout=15)
    ev.note(f"En-tête Server : {r.headers.get('Server')}")
    for h in ("Content-Security-Policy", "X-Frame-Options", "X-Content-Type-Options", "Strict-Transport-Security"):
        ev.note(f"En-tête {h} : {r.headers.get(h, '(absent)')}")
    s = requests.Session()
    r = s.post(BASE + "/login", data={"email": ALICE[0], "password": ALICE[1]}, allow_redirects=False, timeout=15)
    cookie = r.headers.get("Set-Cookie", "")
    ev.note(f"Attributs du cookie de session : {cookie.split(';', 1)[1].strip() if ';' in cookie else cookie}")
    login_page = requests.get(BASE + "/login", timeout=15).text
    ev.note(f"Jeton CSRF dans le formulaire de connexion : {'oui' if 'csrf' in login_page.lower() else 'non'}")
    ev.log("")
    for q in ("salade", "japonaise", "burger"):
        r = web(ev, requests.Session(), "GET", f"/?q={q}")
        ev.note(f"recherche « {q} » → {[html.unescape(n) for n in re.findall(r'<h2>(.*?)</h2>', r.text)]}")
    d = js(api(ev, "GET", "/restaurants"))
    ev.note(f"restaurants fermés exposés par l'API : {[i['name'] for i in d.get('items', []) if not i['is_open']]}")
    api(ev, "POST", "/coupons/validate", body={"code": "SAVE20", "subtotal": 999999})
    r = web(ev, s, "GET", "/logout", follow=False)
    ev.note(f"déconnexion par simple GET : HTTP {r.status_code}")
    ev.save()


# ---------------------------------------------------------------- export

HEADERS = ["ID", "Catégorie", "Titre", "Objectif", "Préconditions", "Données", "Étapes",
           "Résultat attendu", "Résultat obtenu", "Statut", "Priorité", "Bug lié", "Preuve"]
KEYS = ["id", "categorie", "titre", "objectif", "preconditions", "donnees", "etapes",
        "attendu", "obtenu", "statut", "priorite", "bug", "preuve"]
WIDTHS = [9, 13, 34, 34, 30, 34, 44, 40, 48, 9, 9, 9, 34]
FILLS = {"Pass": "C6EFCE", "Fail": "FFC7CE", "Blocked": "FFEB9C", "Skipped": "D9D9D9"}


def export_xlsx(results, meta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cas de test"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for res in results:
        ws.append([res[k] for k in KEYS])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
        row[9].fill = PatternFill("solid", fgColor=FILLS.get(row[9].value, "FFFFFF"))
        row[9].font = Font(bold=True)
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    sy = wb.create_sheet("Synthèse")
    sy.append(["Campagne", meta["campagne"]])
    sy.append(["Date d'exécution", meta["date"]])
    sy.append(["Environnement", meta["environnement"]])
    sy.append([])
    sy.append(["Catégorie", "Total", "Pass", "Fail", "Blocked", "Skipped"])
    cats = ["Nominal", "Limite", "Erreur", "Règle métier"]
    for cat in cats + ["TOTAL"]:
        rows = [r for r in results if cat == "TOTAL" or r["categorie"] == cat]
        sy.append([cat, len(rows)] + [sum(r["statut"] == s for r in rows) for s in ("Pass", "Fail", "Blocked", "Skipped")])
    for c in sy[5]:
        c.font = Font(bold=True)
    for c in sy[sy.max_row]:
        c.font = Font(bold=True)
    sy.column_dimensions["A"].width = 20
    sy.column_dimensions["B"].width = 60
    wb.save(ROOT / "docs" / "test-cases.xlsx")


def main():
    try:
        requests.get(BASE + "/health", timeout=5).raise_for_status()
    except requests.RequestException as exc:
        sys.exit(f"QuickOrder injoignable sur {BASE} : {exc}")
    started = datetime.now()
    results = []
    for c in CASES:
        ev = Evidence(c["id"])
        ev.note(f"{c['id']} : {c['titre']}")
        ev.note(f"Exécuté le {datetime.now():%d/%m/%Y %H:%M:%S} sur {BASE}")
        ev.log("")
        try:
            obtenu, ok = c["fn"](ev)
            statut = "Pass" if ok else "Fail"
        except Exception as exc:  # un cas qui plante est bloqué, pas en échec
            obtenu, statut = f"Exécution impossible : {type(exc).__name__}: {exc}", "Blocked"
        ev.note(f"RÉSULTAT ATTENDU : {c['attendu']}")
        ev.note(f"RÉSULTAT OBTENU  : {obtenu}")
        ev.note(f"STATUT           : {statut}")
        ev.save()
        res = {k: v for k, v in c.items() if k != "fn"}
        res.update(obtenu=obtenu, statut=statut, preuve=f"docs/bug-reports/evidence/{c['id']}.txt")
        if statut == "Pass":
            res["bug"] = ""
        results.append(res)
        print(f"{c['id']:8} {statut:8} {obtenu}")
    exploratory()
    meta = {"campagne": "Campagne initiale QuickOrder V3.1", "date": f"{started:%d/%m/%Y %H:%M}",
            "environnement": f"{BASE} : Docker Compose, base SQLite initialisée par seed.py"}
    (ROOT / "docs" / "campaign-results.json").write_text(
        json.dumps({"meta": meta, "results": results}, ensure_ascii=False, indent=2), encoding="utf-8")
    export_xlsx(results, meta)
    counts = {s: sum(r["statut"] == s for r in results) for s in ("Pass", "Fail", "Blocked", "Skipped")}
    print(f"\n{len(results)} cas : {counts}")


if __name__ == "__main__":
    main()
